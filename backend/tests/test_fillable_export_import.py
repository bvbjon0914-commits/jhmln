# -*- coding: utf-8 -*-
"""
Ausfüllbarer Datenqualität-Export + Direkt-Reimport über den
Zuständigkeiten-Import.

Deckt ab:
- AGS-Normalisierung (führende Nullen, ".0", ungültige Werte) und
  matching_level-Ableitung aus der AGS-Länge,
- Auskunftsart pro Zeile (Name/Code/ID, Groß-/Kleinschreibung, Umlaute, ß,
  Leerraum; unbekannt; leer + Formular-Vorgabe; fehlt ganz),
- nicht ausgefüllte Zeilen -> nur "skipped", kein details-Eintrag,
- Dubletten je Auskunftsart,
- API: 400 ohne Auskunftsart, Zeilen-Auskunftsart per Mapping,
- ROUNDTRIP: Export -> in Excel ausfüllen -> Import -> Lücke geschlossen.
"""
import io
import json

import openpyxl
import pandas as pd
import pytest

from app.api.data_quality import _build_export_sheets, _coverage_gaps, _render_export_xlsx
from app.models.authority import Authority
from app.models.jurisdiction import Jurisdiction
from app.models.settings import AppSettings
from app.services.import_service import ImportService, infer_matching_level, normalize_ags
from tests.conftest import (
    make_authority, make_building, make_jurisdiction, make_request_type,
)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(autouse=True)
def _no_login_required(db_session):
    """imports/data_quality liegen hinter require_login - hier geht es um den
    Import/Export, nicht um den Login (gleiches Muster wie test_matching_api.py)."""
    settings = AppSettings.get_or_create(db_session)
    settings.login_required = False
    db_session.commit()


def df_from_rows(rows):
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# AGS-Normalisierung / Ebenen-Ableitung
# ---------------------------------------------------------------------------

class TestNormalizeAgs:
    @pytest.mark.parametrize("raw,expected", [
        ("05911000", "05911000"),    # 8 Stellen unverändert
        ("5911000", "05911000"),     # 7 -> 8
        ("05911", "05911"),          # 5 unverändert
        ("5911", "05911"),           # 4 -> 5
        ("05", "05"),                # 2 unverändert
        ("5", "05"),                 # 1 -> 2
        ("  05911000  ", "05911000"),
        ("5911000.0", "05911000"),   # Excel-Zahlenformat
        ("5911.0", "05911"),
        ("11000000", "11000000"),    # Berlin: ohne führende Null, 8 Stellen
    ])
    def test_valid(self, raw, expected):
        assert normalize_ags(raw) == expected

    @pytest.mark.parametrize("raw", [
        "", "   ", None, "abc", "0591100A", "059110", "059110000", "123", "123456", "05 911", "-5911000", "5911000.5",
    ])
    def test_invalid(self, raw):
        assert normalize_ags(raw) is None

    def test_int_input(self):
        assert normalize_ags(5911000) == "05911000"


class TestInferMatchingLevel:
    @pytest.mark.parametrize("ags,expected", [
        ("05", "STATE"), ("05911", "COUNTY"), ("05911000", "MUNICIPALITY"),
    ])
    def test_by_length(self, ags, expected):
        assert infer_matching_level(ags) == expected

    def test_other_length_uses_default(self):
        assert infer_matching_level("123", "DISTRICT") == "DISTRICT"
        assert infer_matching_level("123") == "MUNICIPALITY"


# ---------------------------------------------------------------------------
# Importer: Auskunftsart pro Zeile
# ---------------------------------------------------------------------------

@pytest.fixture()
def request_types(db_session):
    make_request_type(db_session, "BAUAKTEN", name="Bauaktenauskunft")
    make_request_type(db_session, "ERSCHLIESSUNG", name="Erschließungsbeiträge / Anliegerbescheinigung")
    make_request_type(db_session, "GRUNDBUCH", name="Grundbuchauskunft")


MAPPING = {"authority_name": "Behörde", "ags": "AGS", "request_type": "Auskunftsart"}


def _import(db_session, rows, mapping=None, request_type_id=None, **kwargs):
    return ImportService(db_session).import_jurisdictions(
        df_from_rows(rows), mapping or MAPPING, request_type_id, **kwargs
    )


def _jurisdictions(db_session):
    db_session.expire_all()
    return db_session.query(Jurisdiction).all()


class TestRowRequestType:
    @pytest.mark.parametrize("cell", [
        "Bauaktenauskunft",
        "BAUAKTEN",
        "bauakten",
        "  bauaktenAUSKUNFT ",
        "Bauakten auskunft",
    ])
    def test_resolves_name_code_and_case(self, db_session, request_types, cell):
        summary = _import(db_session, [{"Behörde": "Bauamt X", "AGS": "05911000", "Auskunftsart": cell}])

        assert summary.imported == 1, summary.details
        assert [j.request_type_id for j in _jurisdictions(db_session)] == ["BAUAKTEN"]

    @pytest.mark.parametrize("cell", [
        "Erschließungsbeiträge / Anliegerbescheinigung",
        "erschliessungsbeitraege / anliegerbescheinigung",
        "ERSCHLIESSUNGSBEITRÄGE/ANLIEGERBESCHEINIGUNG",
        "Erschliessungsbeitrage / Anliegerbescheinigung",   # ä -> a
        "erschliessung",                                     # Code, klein
    ])
    def test_resolves_umlaut_and_eszett_insensitive(self, db_session, request_types, cell):
        summary = _import(db_session, [{"Behörde": "Amt", "AGS": "05911000", "Auskunftsart": cell}])

        assert summary.imported == 1, summary.details
        assert [j.request_type_id for j in _jurisdictions(db_session)] == ["ERSCHLIESSUNG"]

    def test_unknown_request_type_needs_review_never_guessed(self, db_session, request_types):
        summary = _import(db_session, [{"Behörde": "Amt", "AGS": "05911000", "Auskunftsart": "Bauakt"}])

        assert summary.needs_review == 1 and summary.imported == 0
        assert summary.details[0].status == "NEEDS_REVIEW"
        assert summary.details[0].message == "Auskunftsart unbekannt: Bauakt"
        assert _jurisdictions(db_session) == []
        assert db_session.query(Authority).count() == 0  # keine Behörde als Nebenwirkung

    def test_unknown_value_does_not_fall_back_to_form_default(self, db_session, request_types):
        summary = _import(
            db_session,
            [{"Behörde": "Amt", "AGS": "05911000", "Auskunftsart": "Gibtsnicht"}],
            request_type_id="BAUAKTEN",
        )

        assert summary.needs_review == 1
        assert _jurisdictions(db_session) == []

    def test_blank_cell_falls_back_to_form_default(self, db_session, request_types):
        summary = _import(
            db_session,
            [
                {"Behörde": "Amt A", "AGS": "05911000", "Auskunftsart": ""},
                {"Behörde": "Amt B", "AGS": "05913000", "Auskunftsart": "Grundbuchauskunft"},
            ],
            request_type_id="BAUAKTEN",
        )

        assert summary.imported == 2
        by_ags = {j.ags: j.request_type_id for j in _jurisdictions(db_session)}
        assert by_ags == {"05911000": "BAUAKTEN", "05913000": "GRUNDBUCH"}  # Zellwert gewinnt

    def test_blank_cell_without_default_needs_review(self, db_session, request_types):
        summary = _import(db_session, [{"Behörde": "Amt", "AGS": "05911000", "Auskunftsart": ""}])

        assert summary.needs_review == 1 and summary.skipped == 0
        assert summary.details[0].message == "Auskunftsart fehlt"
        assert _jurisdictions(db_session) == []
        assert db_session.query(Authority).count() == 0

    def test_form_default_accepts_name_too(self, db_session, request_types):
        summary = _import(
            db_session, [{"Behörde": "Amt", "AGS": "05911000"}],
            mapping={"authority_name": "Behörde", "ags": "AGS"}, request_type_id="Bauaktenauskunft",
        )

        assert summary.imported == 1
        assert [j.request_type_id for j in _jurisdictions(db_session)] == ["BAUAKTEN"]

    def test_inactive_request_type_resolves_but_active_preferred(self, db_session):
        make_request_type(db_session, "OLD_BAU", name="Bauaktenauskunft", active=False)
        make_request_type(db_session, "NEW_BAU", name="Bauaktenauskunft", active=True)
        make_request_type(db_session, "NUR_ALT", name="Nur Alt", active=False)

        summary = _import(db_session, [
            {"Behörde": "Amt A", "AGS": "05911000", "Auskunftsart": "Bauaktenauskunft"},
            {"Behörde": "Amt B", "AGS": "05913000", "Auskunftsart": "Nur Alt"},
        ])

        assert summary.imported == 2
        by_ags = {j.ags: j.request_type_id for j in _jurisdictions(db_session)}
        assert by_ags == {"05911000": "NEW_BAU", "05913000": "NUR_ALT"}


class TestBackwardCompat:
    def test_form_request_type_id_only_positional(self, db_session, request_types):
        """Altaufrufer: request_type_id als drittes Positionsargument, keine request_type-Spalte."""
        svc = ImportService(db_session)
        df = df_from_rows([{"name": "Bauamt X", "ags_col": "05911000"}])

        summary = svc.import_jurisdictions(df, {"authority_name": "name", "ags": "ags_col"}, "BAUAKTEN")

        assert summary.imported == 1
        (j,) = _jurisdictions(db_session)
        assert (j.request_type_id, j.ags, j.matching_level) == ("BAUAKTEN", "05911000", "MUNICIPALITY")

    def test_unregistered_form_id_still_used_as_is(self, db_session):
        """Bisheriges Verhalten: die Formular-ID wird nicht gegen die Tabelle geprüft."""
        summary = ImportService(db_session).import_jurisdictions(
            df_from_rows([{"name": "Amt", "ags_col": "05911000"}]),
            {"authority_name": "name", "ags": "ags_col"}, "FREMD",
        )

        assert summary.imported == 1
        assert [j.request_type_id for j in _jurisdictions(db_session)] == ["FREMD"]

    def test_missing_ags_with_name_still_needs_review_when_column_not_mapped(self, db_session, request_types):
        summary = ImportService(db_session).import_jurisdictions(
            df_from_rows([{"name": "Bauamt X", "ags_col": ""}]),
            {"authority_name": "name", "ags": "ags_col"}, "BAUAKTEN",
        )

        assert summary.needs_review == 1 and summary.skipped == 0

    def test_summary_to_dict_has_skipped(self, db_session, request_types):
        body = _import(db_session, [{"Behörde": "", "AGS": "", "Auskunftsart": ""}]).to_dict()

        assert body["skipped"] == 1 and body["details"] == []


# ---------------------------------------------------------------------------
# Importer: ausgefüllt vs. nicht ausgefüllt
# ---------------------------------------------------------------------------

class TestSkippedRows:
    def test_blank_authority_name_is_skipped_without_details(self, db_session, request_types):
        summary = _import(db_session, [
            {"Behörde": "", "AGS": "05911000", "Auskunftsart": "BAUAKTEN"},
            {"Behörde": "   ", "AGS": "", "Auskunftsart": ""},
        ], request_type_id="BAUAKTEN")

        assert summary.skipped == 2
        assert summary.total_rows == 2
        assert (summary.imported, summary.needs_review, summary.errors, summary.duplicates) == (0, 0, 0, 0)
        assert summary.details == []

    def test_blank_ags_and_blank_request_type_cell_is_skipped(self, db_session, request_types):
        """Vorausgefüllte Behörde der Export-Seite 'Ohne Zuständigkeit', aber nichts eingetragen."""
        summary = _import(db_session, [{"Behörde": "Amt", "AGS": "", "Auskunftsart": ""}])

        assert summary.skipped == 1
        assert summary.details == []
        assert db_session.query(Authority).count() == 0

    def test_skip_rule_applies_even_with_form_default(self, db_session, request_types):
        summary = _import(db_session, [{"Behörde": "Amt", "AGS": "", "Auskunftsart": ""}], request_type_id="BAUAKTEN")

        assert summary.skipped == 1 and summary.needs_review == 0

    def test_request_type_filled_but_ags_blank_needs_review(self, db_session, request_types):
        summary = _import(db_session, [{"Behörde": "Amt", "AGS": "", "Auskunftsart": "Bauaktenauskunft"}])

        assert summary.needs_review == 1 and summary.skipped == 0
        assert summary.details[0].message == "AGS fehlt"

    def test_name_and_ags_but_blank_request_type_needs_review(self, db_session, request_types):
        summary = _import(db_session, [{"Behörde": "Amt", "AGS": "05911000", "Auskunftsart": ""}])

        assert summary.needs_review == 1 and summary.skipped == 0
        assert summary.details[0].message == "Auskunftsart fehlt"

    def test_mixed_sheet_counts(self, db_session, request_types):
        rows = [{"Behörde": "", "AGS": "05911000", "Auskunftsart": "BAUAKTEN"} for _ in range(50)]
        rows.append({"Behörde": "Amt", "AGS": "5911000", "Auskunftsart": "BAUAKTEN"})
        rows.append({"Behörde": "Amt", "AGS": "abc", "Auskunftsart": "BAUAKTEN"})

        summary = _import(db_session, rows)

        assert summary.skipped == 50
        assert summary.imported == 1
        assert summary.needs_review == 1
        assert len(summary.details) == 2  # nur die tatsächlich bearbeiteten Zeilen


# ---------------------------------------------------------------------------
# Importer: AGS-Normalisierung + Ebene im Lauf
# ---------------------------------------------------------------------------

class TestAgsInImport:
    def test_leading_zero_padded_and_level_inferred(self, db_session, request_types):
        summary = _import(db_session, [
            {"Behörde": "Gemeindeamt", "AGS": "5911000", "Auskunftsart": "BAUAKTEN"},   # 7 -> 8
            {"Behörde": "Kreisamt", "AGS": "5911", "Auskunftsart": "BAUAKTEN"},          # 4 -> 5
            {"Behörde": "Landesamt", "AGS": "5", "Auskunftsart": "BAUAKTEN"},            # 1 -> 2
            {"Behörde": "Zahlenzelle", "AGS": "5913000.0", "Auskunftsart": "BAUAKTEN"},  # Excel .0
        ])

        assert summary.imported == 4
        levels = {j.ags: j.matching_level for j in _jurisdictions(db_session)}
        assert levels == {
            "05911000": "MUNICIPALITY", "05911": "COUNTY", "05": "STATE", "05913000": "MUNICIPALITY",
        }

    def test_explicit_matching_level_wins(self, db_session, request_types):
        mapping = {**MAPPING, "matching_level": "Ebene"}
        _import(db_session, [{"Behörde": "Amt", "AGS": "05911", "Auskunftsart": "BAUAKTEN", "Ebene": "MUNICIPALITY"}],
                mapping=mapping)

        (j,) = _jurisdictions(db_session)
        assert j.matching_level == "MUNICIPALITY"

    def test_blank_explicit_level_cell_falls_back_to_inference(self, db_session, request_types):
        mapping = {**MAPPING, "matching_level": "Ebene"}
        _import(db_session, [{"Behörde": "Amt", "AGS": "05911", "Auskunftsart": "BAUAKTEN", "Ebene": ""}],
                mapping=mapping)

        (j,) = _jurisdictions(db_session)
        assert j.matching_level == "COUNTY"

    @pytest.mark.parametrize("bad", ["abc", "059110", "123", "0591100A", "059110000"])
    def test_invalid_ags_needs_review(self, db_session, request_types, bad):
        summary = _import(db_session, [{"Behörde": "Amt", "AGS": bad, "Auskunftsart": "BAUAKTEN"}])

        assert summary.needs_review == 1 and summary.imported == 0
        assert summary.details[0].message == f"AGS ungültig: {bad}"
        assert db_session.query(Authority).count() == 0

    def test_one_invalid_ags_in_cell_rejects_whole_row_no_partial_import(self, db_session, request_types):
        summary = _import(db_session, [
            {"Behörde": "Amt", "AGS": "05911000, 12x45, 05913000", "Auskunftsart": "BAUAKTEN"},
        ])

        assert summary.needs_review == 1 and summary.imported == 0
        assert "12x45" in summary.details[0].message
        assert _jurisdictions(db_session) == []

    def test_multiple_ags_in_cell_normalized_each(self, db_session, request_types):
        _import(db_session, [{"Behörde": "Amt", "AGS": "5911000; 5913000, 05916000", "Auskunftsart": "BAUAKTEN"}])

        assert {j.ags for j in _jurisdictions(db_session)} == {"05911000", "05913000", "05916000"}


# ---------------------------------------------------------------------------
# Importer: Dubletten je Auskunftsart
# ---------------------------------------------------------------------------

class TestDuplicatesPerRequestType:
    def test_same_authority_and_ags_for_other_request_type_is_not_duplicate(self, db_session, request_types):
        summary = _import(db_session, [
            {"Behörde": "Amt", "AGS": "05911000", "Auskunftsart": "BAUAKTEN", "Ort": "X"},
            {"Behörde": "Amt", "AGS": "05911000", "Auskunftsart": "GRUNDBUCH", "Ort": "X"},
            {"Behörde": "Amt", "AGS": "05911000", "Auskunftsart": "Bauaktenauskunft", "Ort": "X"},
        ], mapping={**MAPPING, "city": "Ort"})

        assert summary.imported == 2
        assert summary.duplicates == 1
        assert {(j.request_type_id, j.ags) for j in _jurisdictions(db_session)} == {
            ("BAUAKTEN", "05911000"), ("GRUNDBUCH", "05911000"),
        }
        assert db_session.query(Authority).count() == 1

    def test_existing_rows_per_request_type_loaded_lazily(self, db_session, request_types):
        authority = make_authority(db_session, name="Amt", city="X")
        make_jurisdiction(db_session, "BAUAKTEN", authority.authority_id, ags="05911000",
                          matching_level="MUNICIPALITY")
        mapping = {**MAPPING, "city": "Ort"}

        summary = _import(db_session, [
            {"Behörde": "Amt", "AGS": "5911000", "Auskunftsart": "BAUAKTEN", "Ort": "X"},    # Dublette (normalisiert)
            {"Behörde": "Amt", "AGS": "5911000", "Auskunftsart": "GRUNDBUCH", "Ort": "X"},   # neu
        ], mapping=mapping)

        assert (summary.imported, summary.duplicates) == (1, 1)


# ---------------------------------------------------------------------------
# API
# ---------------------------------------------------------------------------

def _post_import(app_client, content: bytes, mapping: dict, filename="zust.csv", **data):
    payload = {"mapping": json.dumps(mapping), **data}
    return app_client.post(
        "/api/import/jurisdictions", files={"file": (filename, content, "text/csv")}, data=payload,
    )


CSV_ROWS = (
    "Behörde,AGS,Auskunftsart\n"
    "Bauamt Mitte,5911000,Bauaktenauskunft\n"
    "Grundbuchamt Mitte,05911000,GRUNDBUCH\n"
    ",05913000,BAUAKTEN\n"
).encode("utf-8")


class TestImportApi:
    def test_no_request_type_anywhere_is_400(self, app_client, request_types):
        res = _post_import(app_client, CSV_ROWS, {"authority_name": "Behörde", "ags": "AGS"})

        assert res.status_code == 400
        assert res.json()["detail"] == (
            "Auskunftsart fehlt (Formularfeld request_type_id oder Mapping-Spalte request_type)"
        )

    def test_empty_request_type_mapping_value_is_400(self, app_client, request_types):
        res = _post_import(app_client, CSV_ROWS, {"authority_name": "Behörde", "ags": "AGS", "request_type": ""})

        assert res.status_code == 400

    def test_missing_authority_name_column_is_400(self, app_client, request_types):
        res = _post_import(app_client, CSV_ROWS, {"ags": "AGS", "request_type": "Auskunftsart"})

        assert res.status_code == 400

    def test_per_row_request_type_column(self, app_client, db_session, request_types):
        res = _post_import(app_client, CSV_ROWS, MAPPING)

        assert res.status_code == 200, res.text
        body = res.json()
        assert (body["imported"], body["skipped"], body["needs_review"]) == (2, 1, 0)
        assert body["details"] and all(d["row_index"] in (0, 1) for d in body["details"])
        db_session.expire_all()
        assert {(j.request_type_id, j.ags) for j in db_session.query(Jurisdiction).all()} == {
            ("BAUAKTEN", "05911000"), ("GRUNDBUCH", "05911000"),
        }

    def test_form_request_type_id_only_still_works(self, app_client, db_session, request_types):
        csv_bytes = "Behörde,AGS\nBauamt Mitte,5911000\n".encode("utf-8")
        res = _post_import(app_client, csv_bytes, {"authority_name": "Behörde", "ags": "AGS"},
                           request_type_id="BAUAKTEN")

        assert res.status_code == 200, res.text
        assert res.json()["imported"] == 1
        db_session.expire_all()
        (j,) = db_session.query(Jurisdiction).all()
        assert (j.request_type_id, j.ags, j.matching_level) == ("BAUAKTEN", "05911000", "MUNICIPALITY")

    def test_form_default_plus_request_type_column(self, app_client, db_session, request_types):
        csv_bytes = "Behörde,AGS,Auskunftsart\nAmt A,05911000,\nAmt B,05913000,GRUNDBUCH\n".encode("utf-8")
        res = _post_import(app_client, csv_bytes, MAPPING, request_type_id="BAUAKTEN")

        assert res.status_code == 200, res.text
        assert res.json()["imported"] == 2


# ---------------------------------------------------------------------------
# Export-Layout
# ---------------------------------------------------------------------------

EXPECTED_SHEET_ORDER = [
    "Ohne E-Mail", "Ohne Zuständigkeit", "Ohne Adresse", "Nicht verifiziert", "Behörden-Duplikate",
    "Zuständigkeits-Duplikate", "Verwaiste Zuständigkeiten", "Gebäude-Duplikate", "Gebäude Prüfung nötig",
    "Ohne Kartenkoordinaten", "Abdeckungslücken", "Mögliche Duplikate (ähnlich)", "Anleitung",
]
AUTHORITY_HEADERS = ["Behörde", "Abteilung", "Straße", "Hausnummer", "PLZ", "Ort", "Bundesland", "E-Mail",
                     "Telefon", "Website"]
EXPECTED_HEADERS = {
    "Ohne E-Mail": AUTHORITY_HEADERS,
    "Ohne Zuständigkeit": AUTHORITY_HEADERS + ["Auskunftsart", "AGS", "Gemeinde", "Quelle"],
    "Ohne Adresse": AUTHORITY_HEADERS,
    "Nicht verifiziert": AUTHORITY_HEADERS,
    "Behörden-Duplikate": AUTHORITY_HEADERS,
    "Zuständigkeits-Duplikate": ["Behörde", "Auskunftsart", "AGS", "Gemeinde"],
    "Verwaiste Zuständigkeiten": ["Behörde", "Auskunftsart", "AGS", "Gemeinde"],
    "Gebäude-Duplikate": ["Straße", "Hausnummer", "PLZ", "Ort"],
    "Gebäude Prüfung nötig": ["Straße", "Hausnummer", "PLZ", "Ort"],
    "Ohne Kartenkoordinaten": ["Straße", "Hausnummer", "PLZ", "Ort"],
    "Abdeckungslücken": ["AGS", "Gemeinde", "Auskunftsart", "Betroffene Gebäude", "Behörde", "Abteilung",
                         "Straße", "Hausnummer", "PLZ", "Ort", "E-Mail", "Telefon", "Website", "Quelle"],
    "Mögliche Duplikate (ähnlich)": ["Behörde A", "Behörde B", "Ort", "Ähnlichkeit"],
    "Anleitung": ["Anleitung"],
}


def _sheet_rows(workbook, name):
    return [list(row) for row in workbook[name].iter_rows(values_only=True)]


@pytest.fixture()
def seeded(db_session):
    """Zwei Abdeckungslücken (BAUAKTEN in Köln + Bielefeld, weil dort keine Regel existiert, obwohl es
    anderswo eine gibt) und eine Behörde ohne jede Zuständigkeit."""
    make_request_type(db_session, "BAUAKTEN", name="Bauaktenauskunft")
    make_request_type(db_session, "ALT", name="Veraltete Auskunft", active=False)
    elsewhere = make_authority(db_session, name="Bauamt Bochum", city="Bochum", email="bau@bochum.example")
    make_jurisdiction(db_session, "BAUAKTEN", elsewhere.authority_id, ags="05911000", matching_level="MUNICIPALITY")
    make_building(db_session, street="Domstraße", house_number="1", postal_code="50667", city="Köln",
                  ags="05315000")
    make_building(db_session, street="Alter Markt", house_number="2", postal_code="33602", city="Bielefeld",
                  ags="05711000")
    make_authority(db_session, name="Kreisbauamt Nord", city="Nordstadt", email="info@nord.example",
                   street="Amtsweg", house_number="3", postal_code="12345", authority_id="AUTH-NORD")
    return db_session


class TestExportLayout:
    def test_sheet_order_headers_and_anleitung_last(self, seeded):
        sheets = _build_export_sheets(seeded)

        assert [name for name, _, _ in sheets] == EXPECTED_SHEET_ORDER
        assert {name: columns for name, columns, _ in sheets} == EXPECTED_HEADERS
        assert sheets[-1][0] == "Anleitung"

    def test_rendered_workbook_matches_and_fill_columns_are_empty(self, seeded):
        wb = openpyxl.load_workbook(io.BytesIO(_render_export_xlsx(_build_export_sheets(seeded))))

        assert wb.sheetnames == EXPECTED_SHEET_ORDER
        for name, headers in EXPECTED_HEADERS.items():
            assert [c for c in _sheet_rows(wb, name)[0]] == headers

        # AGS-Spalten im Textformat (Excel schneidet sonst führende Nullen ab)
        assert wb["Abdeckungslücken"].column_dimensions["A"].number_format == "@"
        assert wb["Ohne Zuständigkeit"].column_dimensions["L"].number_format == "@"

        gaps = _sheet_rows(wb, "Abdeckungslücken")[1:]
        assert sorted(r[0] for r in gaps) == ["05315000", "05711000"]   # AGS als Text, mit führender Null
        for row in gaps:
            assert row[1] and row[2] == "Bauaktenauskunft" and row[3] == 1
            assert all(v is None for v in row[4:])

        no_jur = _sheet_rows(wb, "Ohne Zuständigkeit")[1:]
        assert [r[0] for r in no_jur] == ["Kreisbauamt Nord"]
        assert no_jur[0][5] == "Nordstadt" and no_jur[0][7] == "info@nord.example"
        assert all(v is None for v in no_jur[0][10:])

    def test_anleitung_content(self, seeded):
        wb = openpyxl.load_workbook(io.BytesIO(_render_export_xlsx(_build_export_sheets(seeded))))
        lines = [r[0] for r in _sheet_rows(wb, "Anleitung")[1:] if r[0]]
        text = "\n".join(lines)

        assert "Datenimport -> Zuständigkeiten" in text
        assert "Abdeckungslücken" in text and "Ohne Zuständigkeit" in text
        assert "Fehlende Daten ergänzen" in text
        assert "Bauaktenauskunft" in lines                   # aktive Auskunftsart gelistet
        assert "Veraltete Auskunft" not in lines             # inaktive nicht
        assert not any(line.startswith("=") for line in lines)

    def test_endpoint_returns_new_layout(self, app_client, seeded):
        res = app_client.get("/api/data-quality/export-xlsx")

        assert res.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        assert wb.sheetnames[-1] == "Anleitung"


# ---------------------------------------------------------------------------
# ROUNDTRIP: Export -> ausfüllen -> Reimport -> Lücke geschlossen
# ---------------------------------------------------------------------------

# Mapping exakt so, wie es autoMapColumns() im Frontend (frontend/src/types/import.ts,
# JURISDICTION_FIELDS + FIELD_ALIASES) für die Export-Spalten bildet:
FRONTEND_AUTOMAP = {
    "authority_name": "Behörde",
    "ags": "AGS",
    "request_type": "Auskunftsart",
    "city": "Ort",
    "department_name": "Abteilung",
    "street": "Straße",
    "house_number": "Hausnummer",
    "postal_code": "PLZ",
    "state": "Bundesland",          # nur im Blatt "Ohne Zuständigkeit" vorhanden
    "email": "E-Mail",
    "phone": "Telefon",
    "website": "Website",
    "municipality": "Gemeinde",
    "source": "Quelle",
}


def _automap(columns):
    """Nur Felder, deren Spalte im Blatt existiert (wie autoMapColumns)."""
    return {field: col for field, col in FRONTEND_AUTOMAP.items() if col in columns}


def _fill(workbook, sheet_name, row_number, values: dict):
    ws = workbook[sheet_name]
    headers = [c.value for c in ws[1]]
    for header, value in values.items():
        ws.cell(row=row_number, column=headers.index(header) + 1, value=value)


def _to_bytes(workbook) -> bytes:
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def _post_sheet(app_client, content: bytes, sheet: str):
    columns = [c for c in pd.read_excel(io.BytesIO(content), sheet_name=sheet, dtype=str).columns]
    return app_client.post(
        "/api/import/jurisdictions",
        files={"file": ("datenqualitaet_luecken.xlsx", content, XLSX_MIME)},
        data={"mapping": json.dumps(_automap(columns)), "sheet": sheet},
    )


class TestRoundtrip:
    def test_export_fill_reimport_closes_gap(self, app_client, seeded):
        db = seeded
        gaps_before = {g["ags"] for g in _coverage_gaps(db)}
        assert gaps_before == {"05315000", "05711000"}

        # --- Export (Endpunkt) ---
        res = app_client.get("/api/data-quality/export-xlsx")
        assert res.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(res.content))

        # --- Ausfüllen: Köln-Lücke komplett, Bielefeld bleibt leer ---
        gap_rows = _sheet_rows(wb, "Abdeckungslücken")
        koeln_row = next(i for i, r in enumerate(gap_rows, start=1) if r[0] == "05315000")
        _fill(wb, "Abdeckungslücken", koeln_row, {
            "Behörde": "Bauaktenarchiv Köln", "Ort": "Köln", "E-Mail": "bauakten@koeln.example",
            "Straße": "Archivstraße", "Hausnummer": "7", "PLZ": "50667", "Quelle": "https://www.stadt-koeln.example/bauakten",
        })
        # --- "Ohne Zuständigkeit": Kreis OHNE führende Null (4 Ziffern, als Zahl wie in Excel) ---
        nord_row = next(i for i, r in enumerate(_sheet_rows(wb, "Ohne Zuständigkeit"), start=1)
                        if r[0] == "Kreisbauamt Nord")
        _fill(wb, "Ohne Zuständigkeit", nord_row, {
            "Auskunftsart": "bauaktenauskunft", "AGS": 5558, "Quelle": "Amtsblatt 2024",
        })
        content = _to_bytes(wb)

        # --- Reimport Blatt 1 ---
        res1 = _post_sheet(app_client, content, "Abdeckungslücken")
        assert res1.status_code == 200, res1.text
        body1 = res1.json()
        assert (body1["imported"], body1["needs_review"], body1["errors"], body1["duplicates"]) == (1, 0, 0, 0)
        assert body1["skipped"] == 1                       # die unausgefüllte Bielefeld-Zeile
        assert len(body1["details"]) == 1                  # ... ohne details-Eintrag
        assert body1["total_rows"] == 2

        db.expire_all()
        koeln_auth = db.query(Authority).filter_by(authority_name="Bauaktenarchiv Köln", city="Köln").one()
        assert koeln_auth.email == "bauakten@koeln.example"
        assert (koeln_auth.street, koeln_auth.house_number, koeln_auth.postal_code) == ("Archivstraße", "7", "50667")
        assert koeln_auth.source == "https://www.stadt-koeln.example/bauakten"
        (j1,) = db.query(Jurisdiction).filter_by(authority_id=koeln_auth.authority_id).all()
        assert (j1.request_type_id, j1.ags, j1.matching_level) == ("BAUAKTEN", "05315000", "MUNICIPALITY")
        assert j1.municipality == "Köln"
        assert j1.source == "https://www.stadt-koeln.example/bauakten"

        # --- Reimport Blatt 2 ---
        res2 = _post_sheet(app_client, content, "Ohne Zuständigkeit")
        assert res2.status_code == 200, res2.text
        body2 = res2.json()
        assert (body2["imported"], body2["needs_review"], body2["errors"]) == (1, 0, 0)
        assert body2["total_rows"] == 1 and body2["skipped"] == 0

        db.expire_all()
        nord = db.query(Authority).filter_by(authority_id="AUTH-NORD").one()
        (j2,) = db.query(Jurisdiction).filter_by(authority_id="AUTH-NORD").all()
        assert (j2.request_type_id, j2.ags, j2.matching_level) == ("BAUAKTEN", "05558", "COUNTY")
        assert j2.source == "Amtsblatt 2024"
        assert db.query(Authority).filter_by(authority_name="Kreisbauamt Nord").count() == 1  # nicht neu angelegt
        assert nord.email == "info@nord.example"

        # --- zweiter, identischer Import: nur noch Dubletten ---
        total_before = db.query(Jurisdiction).count()
        again1 = _post_sheet(app_client, content, "Abdeckungslücken").json()
        again2 = _post_sheet(app_client, content, "Ohne Zuständigkeit").json()
        assert (again1["imported"], again1["duplicates"], again1["skipped"]) == (0, 1, 1)
        assert (again2["imported"], again2["duplicates"]) == (0, 1)
        db.expire_all()
        assert db.query(Jurisdiction).count() == total_before

        # --- Lücke geschlossen, die nicht ausgefüllte bleibt ---
        gaps_after = {g["ags"] for g in _coverage_gaps(db)}
        assert gaps_after == {"05711000"}

    def test_unfilled_export_imports_to_nothing(self, app_client, seeded):
        """Wer den Export unverändert zurückspielt, erzeugt weder Daten noch Prüffälle."""
        content = app_client.get("/api/data-quality/export-xlsx").content

        body = _post_sheet(app_client, content, "Abdeckungslücken").json()
        assert (body["imported"], body["needs_review"], body["errors"]) == (0, 0, 0)
        assert body["skipped"] == body["total_rows"] == 2 and body["details"] == []

        body = _post_sheet(app_client, content, "Ohne Zuständigkeit").json()
        assert (body["imported"], body["needs_review"], body["errors"]) == (0, 0, 0)
        assert body["skipped"] == body["total_rows"] == 1 and body["details"] == []

    def test_anleitung_is_never_default_sheet(self, seeded):
        content = _render_export_xlsx(_build_export_sheets(seeded))
        sheets = ImportService.list_sheets(content, "x.xlsx")

        assert sheets[-1]["name"] == "Anleitung"
        assert ImportService.pick_default_sheet(sheets) != "Anleitung"
