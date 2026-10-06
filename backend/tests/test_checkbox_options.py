# -*- coding: utf-8 -*-
"""
Ankreuzbare Optionen im Anschreiben: Optionslisten (Single Source of Truth),
GET /api/request-types (checkbox_options), POST /api/matching
(checkbox_selections -> RequestItem.selected_options) und das Rendern der
Auswahl als ☒/☐ im generierten Word-Dokument.
"""
import os
import re

import pytest
from docx import Document

from app.config import TEMPLATES_DIR
from app.models.request import RequestItem
from app.models.settings import AppSettings
from app.services.document_generator import DocumentGenerationService
from app.services.request_options import CHECKBOX_OPTIONS, get_checkbox_options, sanitize_selection
from tests.conftest import (
    make_authority, make_building, make_request, make_request_item, make_request_type,
)

CHECKED = "☒"
UNCHECKED = "☐"


@pytest.fixture(autouse=True)
def _no_login_required(db_session):
    """Die Router liegen hinter require_login - hier geht es nicht um den Login."""
    settings = AppSettings.get_or_create(db_session)
    settings.login_required = False
    db_session.commit()


def _option_paragraphs(doc, request_code):
    """Die Absaetze der Optionen (Text beginnt mit ☒/☐ + Optionslabel)."""
    options = get_checkbox_options(request_code)
    paragraphs = []
    for option in options:
        matches = [p.text for p in doc.paragraphs if p.text.endswith(option)]
        assert len(matches) == 1, f"Option nicht genau einmal im Dokument: {option!r}"
        paragraphs.append(matches[0])
    return paragraphs


class TestOptionLists:
    def test_options_match_the_eleven_template_keys(self):
        template_codes = {f[:-5] for f in os.listdir(TEMPLATES_DIR) if f.endswith(".docx")}
        assert len(CHECKBOX_OPTIONS) == 11
        assert set(CHECKBOX_OPTIONS) == template_codes

    def test_every_auskunftsart_has_options_and_no_empty_label(self):
        for code, options in CHECKBOX_OPTIONS.items():
            assert options, code
            assert all(o.strip() for o in options), code

    def test_get_checkbox_options_is_case_insensitive_and_defensive(self):
        assert get_checkbox_options("BAUAKTEN") == CHECKBOX_OPTIONS["bauakten"]
        assert get_checkbox_options("bauakten") == CHECKBOX_OPTIONS["bauakten"]
        assert get_checkbox_options("UNBEKANNT") == []
        assert get_checkbox_options(None) == []
        assert get_checkbox_options("") == []

    def test_get_checkbox_options_returns_a_copy(self):
        get_checkbox_options("GRUNDBUCH").append("manipuliert")
        assert "manipuliert" not in CHECKBOX_OPTIONS["grundbuch"]

    def test_sanitize_selection_drops_invalid_indices(self):
        # BAUAKTEN hat 3 Optionen
        assert sanitize_selection("BAUAKTEN", [2, 0, 0, -1, 3, 99]) == [0, 2]
        assert sanitize_selection("BAUAKTEN", [True, "1", 1.5, None, 1]) == [1]
        assert sanitize_selection("BAUAKTEN", None) == []
        assert sanitize_selection("UNBEKANNT", [0, 1]) == []


class TestRequestTypesEndpoint:
    def test_request_types_include_checkbox_options(self, app_client, db_session):
        make_request_type(db_session, "BAUAKTEN")
        make_request_type(db_session, "GRUNDBUCH")

        response = app_client.get("/api/request-types")

        assert response.status_code == 200
        by_id = {rt["request_type_id"]: rt for rt in response.json()}
        assert by_id["BAUAKTEN"]["checkbox_options"] == CHECKBOX_OPTIONS["bauakten"]
        assert len(by_id["BAUAKTEN"]["checkbox_options"]) == 3
        assert by_id["GRUNDBUCH"]["checkbox_options"] == CHECKBOX_OPTIONS["grundbuch"]

    def test_unknown_request_type_gets_empty_list(self, app_client, db_session):
        make_request_type(db_session, "SONDERFALL")

        response = app_client.get("/api/request-types")

        assert response.status_code == 200
        assert response.json()[0]["checkbox_options"] == []


class TestMatchingStoresSelection:
    def _post(self, app_client, db_session, selections):
        make_request_type(db_session, "BAUAKTEN")
        make_request_type(db_session, "GRUNDBUCH")
        building = make_building(db_session)
        payload = {
            "building_id": building.building_id,
            "request_type_ids": ["BAUAKTEN", "GRUNDBUCH"],
        }
        if selections is not None:
            payload["checkbox_selections"] = selections
        response = app_client.post("/api/matching", json=payload)
        assert response.status_code == 200, response.text
        items = db_session.query(RequestItem).all()
        return {i.request_type_id: i for i in items}

    def test_selection_is_stored_on_items(self, app_client, db_session):
        items = self._post(app_client, db_session, {"BAUAKTEN": [0, 2], "GRUNDBUCH": [1]})

        assert items["BAUAKTEN"].selected_options == "[0, 2]"
        assert items["BAUAKTEN"].get_selected_options_list() == [0, 2]
        assert items["GRUNDBUCH"].get_selected_options_list() == [1]

    def test_without_selection_nothing_is_stored(self, app_client, db_session):
        items = self._post(app_client, db_session, None)

        assert items["BAUAKTEN"].selected_options is None
        assert items["BAUAKTEN"].get_selected_options_list() == []

    def test_invalid_input_is_dropped_never_an_error(self, app_client, db_session):
        items = self._post(
            app_client,
            db_session,
            {
                "BAUAKTEN": [2, 2, -1, 3, 17, 0],  # Duplikat, negativ, ausserhalb
                "GRUNDBUCH": [5],  # nur ungueltige -> nichts angekreuzt
                "NICHT_ANGEFRAGT": [0],  # nicht angefragte/unbekannte Art
            },
        )

        assert items["BAUAKTEN"].get_selected_options_list() == [0, 2]
        assert items["GRUNDBUCH"].selected_options is None
        assert len(items) == 2


class TestGeneratedDocumentCheckboxes:
    def _generate(self, db_session, tmp_path, code, selected=None):
        service = DocumentGenerationService(templates_dir=TEMPLATES_DIR, output_dir=str(tmp_path))
        building = make_building(db_session)
        authority = make_authority(db_session)
        request_type = make_request_type(db_session, code)
        kwargs = {} if selected is None else {"selected_options": selected}
        result = service.generate_document(building, authority, request_type, "AZ-1", **kwargs)
        return Document(result.filepath)

    def test_ticked_option_is_checked_and_others_unchecked(self, db_session, tmp_path):
        doc = self._generate(db_session, tmp_path, "BAUAKTEN", selected=[1])

        paragraphs = _option_paragraphs(doc, "BAUAKTEN")
        assert paragraphs[0].startswith(UNCHECKED + " ")
        assert paragraphs[1].startswith(CHECKED + " ")
        assert paragraphs[2].startswith(UNCHECKED + " ")

    def test_multiple_ticked_options(self, db_session, tmp_path):
        doc = self._generate(db_session, tmp_path, "BAUAKTEN", selected=[0, 2])

        paragraphs = _option_paragraphs(doc, "BAUAKTEN")
        assert [p[0] for p in paragraphs] == [CHECKED, UNCHECKED, CHECKED]

    def test_default_without_selection_is_all_unchecked(self, db_session, tmp_path):
        doc = self._generate(db_session, tmp_path, "GRUNDBUCH")

        paragraphs = _option_paragraphs(doc, "GRUNDBUCH")
        assert len(paragraphs) == 2
        assert all(p.startswith(UNCHECKED + " ") for p in paragraphs)
        assert CHECKED not in "\n".join(p.text for p in doc.paragraphs)

    def test_out_of_range_indices_are_ignored(self, db_session, tmp_path):
        doc = self._generate(db_session, tmp_path, "GRUNDBUCH", selected=[-1, 7, 1])

        paragraphs = _option_paragraphs(doc, "GRUNDBUCH")
        assert [p[0] for p in paragraphs] == [UNCHECKED, CHECKED]

    @pytest.mark.parametrize("code", sorted(CHECKBOX_OPTIONS))
    def test_no_placeholder_left_over_in_any_template(self, db_session, tmp_path, code):
        n = len(CHECKBOX_OPTIONS[code])
        doc = self._generate(db_session, tmp_path, code.upper(), selected=[0])

        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert not re.search(r"\{\{.*?\}\}", full_text), full_text
        assert "{%" not in full_text
        paragraphs = _option_paragraphs(doc, code)
        assert len(paragraphs) == n
        assert paragraphs[0].startswith(CHECKED + " ")
        assert all(p.startswith(UNCHECKED + " ") for p in paragraphs[1:])

    def test_build_context_has_one_key_per_option(self, db_session, tmp_path):
        service = DocumentGenerationService(templates_dir=TEMPLATES_DIR, output_dir=str(tmp_path))
        request_type = make_request_type(db_session, "BAUAKTEN")
        context = service.build_context(
            make_building(db_session), make_authority(db_session), request_type, "AZ-1",
            selected_options=[2],
        )

        assert [context[f"checkbox_{i}"] for i in range(3)] == [UNCHECKED, UNCHECKED, CHECKED]
        assert "checkbox_3" not in context


class TestGenerationEndpointsUseStoredSelection:
    def test_documents_generate_renders_stored_selection(self, app_client, db_session):
        make_request_type(db_session, "BAUAKTEN")
        authority = make_authority(db_session)
        building = make_building(db_session)
        request = make_request(db_session, building.building_id)
        item = make_request_item(
            db_session, request.request_id, "BAUAKTEN",
            matching_status="MATCHED", authority_id=authority.authority_id,
        )
        item.set_selected_options([2])
        db_session.commit()

        response = app_client.post("/api/documents/generate", json={"request_id": request.request_id})

        assert response.status_code == 200, response.text
        documents = response.json()["documents"]
        assert len(documents) == 1
        doc = Document(documents[0]["filepath"])
        paragraphs = _option_paragraphs(doc, "BAUAKTEN")
        assert [p[0] for p in paragraphs] == [UNCHECKED, UNCHECKED, CHECKED]
